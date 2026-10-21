"""Real Chromium DOM/UI tests with local fixtures. No supplier requests or messages.

The cloud Chromium policy blocks unpacked extensions, so the content script's
message transport and dashboard's Chrome API are replaced by explicit test doubles.
The actual core, content script, dashboard and file writer run unmodified.
"""
import functools
import io
import json
from openpyxl import load_workbook
import http.server
import os
import re
from pathlib import Path
import threading
import time
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / 'test-results'
RESULTS.mkdir(exist_ok=True)
checks = 0

def passed(name):
    global checks
    checks += 1
    print(f'PASS {name}', flush=True)

class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass

handler = functools.partial(QuietHandler, directory=str(ROOT))
server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
threading.Thread(target=server.serve_forever, daemon=True).start()
BASE = f'http://127.0.0.1:{server.server_port}'
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH', '/usr/bin/chromium'), headless=True, ignore_default_args=['--disable-popup-blocking'])
        context = browser.new_context(viewport={'width': 1440, 'height': 1050}, timezone_id='Asia/Irkutsk')
        errors = []
        context.on('page', lambda page: page.on('pageerror', lambda error: errors.append(str(error))))
        ui = context.new_page()
        def dashboard_route(route):
            html = (ROOT/'extension/dashboard.html').read_text()
            html = html.replace('<script src="dashboard.js" type="module">', '<script src="/tests/ui-harness.js" type="module"></script><script src="dashboard.js" type="module">')
            route.fulfill(body=html, content_type='text/html')
        context.route('**/extension/dashboard.html', dashboard_route)
        ui.goto(BASE+'/extension/dashboard.html')
        expect(ui.locator('#queue-count')).to_have_text('0')
        ui.get_by_role('button', name='Сообщение и настройки').click()
        ui.locator('#message').fill('Hello! Please contact me on WeChat: demo-buyer')
        ui.get_by_role('button', name='Сохранить настройки').click()
        expect(ui.locator('#toast')).to_have_text('Настройки сохранены')
        ui.get_by_role('button', name='Очередь поставщиков').click()
        ui.locator('#examples').click()
        ui.wait_for_function("document.querySelector('#supplier-input').value.includes('yiyueda')")
        ui.locator('#add').click()
        expect(ui.locator('#queue-count')).to_have_text('12')
        expect(ui.locator('#queue-body tr')).to_have_count(12)
        passed('dashboard saves message and imports the 12 supplied stores')
        ui.locator('#supplier-input').fill('https://jcnyjxsb.en.alibaba.com/')
        ui.locator('#add').click()
        expect(ui.locator('#import-result')).to_contain_text('уже есть')
        expect(ui.locator('#queue-count')).to_have_text('12')
        passed('duplicate import is rejected by the UI')
        ui.locator('#start').click()
        expect(ui.locator('#run-state')).to_have_text('Очередь работает')
        ui.locator('#pause').click()
        expect(ui.locator('#run-state')).to_have_text('На паузе')
        passed('start and pause update dashboard state')
        ui.locator('#connect-folder').click()
        expect(ui.locator('#folder-status')).to_contain_text('сохранено')
        file_data = ui.evaluate("""async()=>{
          const directory=await navigator.storage.getDirectory();
          const read=async name=>await (await (await directory.getFileHandle(name)).getFile()).text();
          return {schedule:await read('schedule.txt'),report:await read('report.txt'),log:await read('activity.log'),excel:Array.from(new Uint8Array(await (await (await directory.getFileHandle('report.xlsx')).getFile()).arrayBuffer()))};
        }""")
        assert 'jcnyjxsb.en.alibaba.com' in file_data['schedule']
        assert 'Browser UI test' in file_data['report']
        assert 'Импорт' in file_data['log']
        assert 'Asia/Irkutsk' in file_data['report'] and '+08:00' in file_data['log']
        automatic_book=load_workbook(io.BytesIO(bytes(file_data['excel'])))
        assert automatic_book.sheetnames==['Результаты','События']
        assert automatic_book['Результаты'].max_row==13
        assert 'Asia/Irkutsk' in automatic_book['Результаты']['G2'].value
        passed('auto-write creates Excel, TXT and LOG reports with the computer timezone using File System Access')
        protection=ui.evaluate("""async()=>{
          const {writeLog,activityText}= {...await import('/extension/lib/files.js'),...await import('/extension/lib/core.js')};
          const folder=await navigator.storage.getDirectory();
          const handle=await folder.getFileHandle('activity.log');
          const writer=await handle.createWritable();await writer.write('My unrelated log');await writer.close();
          let rejected=false;try{await writeLog(folder,'# OpenLinks — activity.log\\nnew log');}catch{rejected=true;}
          const preserved=await (await handle.getFile()).text();
          await writeLog(folder,'# OpenLinks — activity.log\\nRestored browser test log',true);
          return {rejected,preserved};
        }""")
        assert protection=={'rejected':True,'preserved':'My unrelated log'}
        passed('new activity.log cannot silently overwrite an unrelated existing file')
        excel_protection=ui.evaluate("""async()=>{
          const {writeExcel}=await import('/extension/lib/files.js'),{reportExcel}=await import('/extension/lib/excel.js'),{initialState}=await import('/extension/lib/core.js');
          const folder=await navigator.storage.getDirectory(),bytes=reportExcel(initialState());
          await writeExcel(folder,bytes);
          const handle=await folder.getFileHandle('report.xlsx'),writer=await handle.createWritable();await writer.write('Unrelated workbook');await writer.close();
          let rejected=false;try{await writeExcel(folder,bytes);}catch{rejected=true;}
          const preserved=await (await handle.getFile()).text();await writeExcel(folder,bytes,true);
          return {rejected,preserved};
        }""")
        assert excel_protection=={'rejected':True,'preserved':'Unrelated workbook'}
        passed('Excel auto-write replaces its own workbook and preserves an unrelated existing workbook')

        with ui.expect_download() as download:
            ui.locator('#export-schedule').click()
        assert download.value.suggested_filename == 'schedule.txt'
        passed('manual TXT download')
        ui.locator('#show-log').click()
        ui.locator('#log-filter').fill('неизвестное-событие')
        expect(ui.locator('#report-body tr')).to_have_count(0)
        ui.locator('#log-filter').fill('импорт')
        expect(ui.locator('#report-body tr')).to_have_count(2)
        with ui.expect_download() as log_download:
            ui.locator('#log-download').click()
        assert log_download.value.suggested_filename=='activity.log'
        ui.get_by_role('button', name='Очередь поставщиков').click()
        passed('detailed live log has search and downloadable activity.log')
        ui.screenshot(path=str(RESULTS/'dashboard.png'), full_page=True)
        ui.set_viewport_size({'width': 600, 'height': 850})
        assert ui.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        passed('responsive layout at 600px')
        ui.set_viewport_size({'width':1440,'height':1050})
        ui.get_by_role('button',name='Сообщение и настройки').click()
        ui.evaluate("seedChatRoute('https://example-shop.1688.com/','https://air.1688.com/app/ocms-fusion-components-1688/def_cbu_web_im/index.html?touid=cnalichnexample&siteid=cnalichn&status=1#/')")
        expect(ui.locator('#known-chats')).to_contain_text('example-shop.1688.com')
        expect(ui.locator('#known-chats a')).to_have_attribute('href',re.compile('touid=cnalichnexample'))
        assert ui.locator('#known-chats button').get_attribute('type')=='button'
        ui.locator('#known-chats button').click()
        expect(ui.locator('#known-chats a')).to_have_count(0)
        ui.get_by_role('button',name='Очередь поставщиков').click()
        passed('dashboard lists saved 1688 chat associations and forgets them without submitting settings')
        ui.evaluate('seedLinkResults()')
        ui.locator('#show-log').click()
        expect(ui.locator('#results-body tr')).to_have_count(4)
        expect(ui.locator('#result-summary')).to_contain_text('Ошибки / нужна проверка: 2')
        ui.locator('#result-group').select_option('errors')
        expect(ui.locator('#results-body tr')).to_have_count(2)
        ui.locator('#result-search').fill('same.alibaba.com')
        expect(ui.locator('#results-body tr')).to_have_count(1)
        ui.locator('#results-body button').click()
        expect(ui.locator('#report-body tr')).to_have_count(1)
        expect(ui.locator('#report-body')).to_contain_text('Three errors')
        ui.locator('#all-events').click()
        expect(ui.locator('#report-body tr')).to_have_count(3)
        ui.locator('#result-group').select_option('all')
        ui.locator('#result-search').fill('')
        with ui.expect_download() as report_download:
            ui.locator('#report-download').click()
        report=Path(report_download.value.path()).read_text()
        assert 'РЕЗУЛЬТАТЫ ПО ССЫЛКАМ' in report and 'failed\t3' in report and 'СОБЫТИЯ' in report
        expect(ui.locator('#report-timezone')).to_contain_text('Asia/Irkutsk')
        with ui.expect_download() as excel_download:
            ui.locator('#excel-download').click()
        assert excel_download.value.suggested_filename=='report.xlsx'
        book=load_workbook(io.BytesIO(Path(excel_download.value.path()).read_bytes()))
        assert book['Результаты'].max_row==5 and book['События'].max_row==4
        assert book['Результаты'].freeze_panes=='A2'
        for row in list(book['Результаты'].rows)[1:]:
            assert row[5].is_date and row[6].value=='Asia/Irkutsk (UTC+08:00)'
            assert ui.locator('#results-body').inner_text().find(row[5].value.strftime('%d.%m.%Y'))>=0
            assert row[5].value.strftime('%H:%M:%S') in ui.locator('#results-body').inner_text()
        passed('Excel download contains all link results and events with real local dates matching the dashboard')

        ui.screenshot(path=str(RESULTS/'link-journal.png'),full_page=True)
        ui.set_viewport_size({'width':600,'height':850})
        assert ui.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        ui.set_viewport_size({'width':1440,'height':1050})
        passed('per-link journal separates repeated URLs, filters errors, opens task events and downloads a full report')
        ui.evaluate('restoreLinkResults()')
        ui.get_by_role('button',name='Очередь поставщиков').click()


        schedule_ui=context.new_page()
        schedule_ui.add_init_script('Date.now=()=>1900000000000;')
        schedule_ui.goto(BASE+'/extension/dashboard.html')
        schedule_ui.get_by_role('button',name='Сообщение и настройки').click()
        schedule_ui.locator('#interval-1688').fill('137')
        schedule_ui.get_by_role('button',name='Сохранить настройки').click()
        expect(schedule_ui.locator('#toast')).to_have_text('Настройки сохранены')
        schedule_ui.get_by_role('button',name='Очередь поставщиков').click()
        schedule_ui.locator('#supplier-input').fill('\n'.join(
            link for index in range(5) for link in [f'https://fixture{index}.1688.com/',f'https://fixture{index}.en.made-in-china.com/']))
        schedule_ui.locator('#add').click()
        expect(schedule_ui.locator('#queue-body tr')).to_have_count(10)
        forecast=schedule_ui.evaluate("""async()=>{
          const {state}=await chrome.runtime.sendMessage({type:'GET'});
          const {makePlan,scheduleText}=await import('/extension/lib/core.js');
          const {displayTime}=await import('/extension/lib/time.js');
          const dates=makePlan(state,Date.now()).filter(row=>row.platform==='1688').map(row=>row.plannedAt);
          const displayed=[...document.querySelectorAll('#queue-body tr')].filter(row=>row.children[2].textContent==='1688').map(row=>row.children[3].textContent);
          return {dates,expected:dates.map(displayTime),displayed,schedule:scheduleText(state,Date.now())};
        }""")
        assert forecast['dates']==[1900000000000+index*137000 for index in range(5)]
        assert forecast['displayed']==forecast['expected'] and len(set(forecast['displayed']))==5
        schedule_lines=[line.split('\t')[0] for line in forecast['schedule'].splitlines() if '\t1688\t' in line]
        assert len(set(schedule_lines))==5
        schedule_ui.close()
        passed('dashboard and exported schedule show five distinct 1688 times 137 seconds apart despite intervening MIC shops')

        context.route('https://test.en.alibaba.com/**', lambda route: route.fulfill(
            body=(ROOT/'tests/fixtures'/('chat.html' if '/chat' in route.request.url else 'store.html')).read_text(), content_type='text/html'))
        supplier = context.new_page()
        contact_source=(ROOT/'extension/lib/native-contact.js').read_text().replace('export ','')
        adapter_transport=contact_source+"""
          let listener;
          window.lessonMessages=[];
          window.adapterMessages=[];
          window.sendGuardMode='allow';
          window.chrome={runtime:{id:'adapter-test',onMessage:{addListener(fn){listener=fn}},sendMessage:async message=>{adapterMessages.push(message);if(message.type==='OPENLINKS_TEST_SEND_GUARD'){if(sendGuardMode==='error')throw Error('Guard transport unavailable');return sendGuardMode==='allow'?{ok:true,canSend:true}:{ok:true,canSend:false,reasonCode:'document_mismatch'};}if(message.type==='LEARN_PICK')lessonMessages.push(message);return {ok:true};}}};
          window.contactCommands=[];
          chrome.debugger={attach:async()=>{},detach:async()=>{},sendCommand:(target,method,params)=>new Promise(resolve=>contactCommands.push({method,params,resolve}))};
          const rawAdapter=(type,extra={})=>new Promise(resolve=>listener({type,jobId:'test-job',profile:{},message:'Hello, WeChat: demo',...extra},{id:'adapter-test'},resolve));
          window.adapter=async(type,extra={})=>{
            const result=await rawAdapter(type,extra);
            return type==='CONTACT'&&result.ok?clickNativeContact({tabId:1},result):result;
          };
        """
        contact_cdp=context.new_cdp_session(supplier)
        def contact_eval(expression):
            return contact_cdp.send('Runtime.evaluate',{'expression':expression,'returnByValue':True,'userGesture':False})['result'].get('value')
        def contact_call(expression):
            contact_eval('window.contactResult=null;Promise.resolve('+expression+').then(value=>{contactResult={value}},error=>{contactResult={error:error.message}});void 0')
            deadline=time.monotonic()+10
            while time.monotonic()<deadline:
                command=contact_eval('contactCommands[0]&&{method:contactCommands[0].method,params:contactCommands[0].params}')
                if command:
                    reply=contact_cdp.send(command['method'],command['params'])
                    contact_eval('contactCommands.shift().resolve('+json.dumps(reply)+');void 0')
                result=contact_eval('contactResult')
                if result:
                    assert 'error' not in result,result
                    return result['value']
                supplier.wait_for_timeout(1)
            raise AssertionError('Native contact CDP bridge timed out')
        supplier.add_init_script(adapter_transport)
        supplier.goto('https://test.en.alibaba.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        result = supplier.evaluate("adapter('PROBE')")
        assert result['kind']=='contact' and result['href'] is None
        passed('semantic contact discovery')
        supplier.goto('https://test.en.alibaba.com/chat')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        assert supplier.evaluate("adapter('PROBE')")['kind']=='chat'
        result = supplier.evaluate("adapter('PREPARE')")
        assert result['ok']
        expect(supplier.locator('#message')).to_have_value('Hello, WeChat: demo')
        passed('preparation fills React-style native input and enables initially disabled Send')
        supplier.evaluate("sendGuardMode='deny'")
        denied=supplier.evaluate("adapter('SEND', {attemptId:'guard-denied',expectedUrl:location.href})")
        assert not denied['clicked'] and denied['notAttempted'] and supplier.evaluate('clicks')==0
        assert denied['permissionReasonCode']=='document_mismatch' and 'document_mismatch' in denied['reason']
        expect(supplier.locator('#message')).to_have_value('Hello, WeChat: demo')
        supplier.evaluate("sendGuardMode='error'")
        disconnected=supplier.evaluate("adapter('SEND', {attemptId:'guard-unavailable',expectedUrl:location.href})")
        assert not disconnected['clicked'] and disconnected['notAttempted'] and supplier.evaluate('clicks')==0
        supplier.evaluate("sendGuardMode='allow'")
        passed('permission denial and guard transport failure prove no click and preserve the prepared draft')
        assert supplier.evaluate("adapter('SEND', {expectedUrl:location.href})")['clicked']
        assert supplier.evaluate('clicks')==1
        assert not supplier.evaluate("adapter('SEND', {expectedUrl:location.href})")['clicked']
        supplier.evaluate("sendGuardMode='deny'")
        duplicate=supplier.evaluate("adapter('SEND', {expectedUrl:location.href})")
        assert not duplicate['clicked'] and duplicate['notAttempted'] is False and supplier.evaluate('clicks')==1
        supplier.evaluate("sendGuardMode='allow'")
        passed('a denied duplicate of a previously clicked operation never authorizes releasing its quota')
        passed('single send without result checks and duplicate-click prevention')
        supplier.evaluate("()=>{let e=document.querySelector('#message');e.value='My existing draft';window.replacementInputs=[];e.addEventListener('input',()=>replacementInputs.push(e.value));}")
        replaced=supplier.evaluate("adapter('PREPARE')")
        assert replaced['ok'] and replaced['draftReplaced']
        expect(supplier.locator('#message')).to_have_value('Hello, WeChat: demo')
        assert supplier.evaluate('replacementInputs')==['','Hello, WeChat: demo']
        assert supplier.evaluate('clicks')==1
        passed('Alibaba clears the old draft before inserting the new message without prematurely sending')
        supplier.evaluate("document.body.insertAdjacentHTML('beforeend','<div id=fixture-captcha style=width:200px;height:60px>Check</div>')")
        assert 'CAPTCHA' in supplier.evaluate("adapter('PROBE')")['obstacle']
        assert not supplier.evaluate("adapter('SEND', {jobId:'another',expectedUrl:location.href})")['clicked']
        passed('visible CAPTCHA blocks preparation and sending')
        supplier.evaluate("document.querySelector('#fixture-captcha').remove();document.querySelector('#chat-panel').insertAdjacentHTML('beforeend','<textarea id=second></textarea>')")
        assert supplier.evaluate("adapter('PROBE')")['kind']=='ambiguous'
        passed('ambiguous editors stop automation')
        assert supplier.evaluate("adapter('PROBE',{profile:{editor:'['}})")['kind']=='ambiguous'
        passed('invalid CSS selector fails closed')
        supplier.evaluate("document.querySelector('#second').remove();document.querySelector('#message').value='';document.querySelector('form').insertAdjacentHTML('afterbegin','<input required type=email id=email>')")
        assert 'запрос' in supplier.evaluate("adapter('PREPARE')")['reason']
        passed('inquiry-like form with additional email field is rejected')
        supplier.evaluate("document.body.innerHTML='<form id=sendInquiry><textarea id=message></textarea><button id=send>Send inquiry</button></form>'")
        assert supplier.evaluate("adapter('PROBE')")['kind']=='inquiry'
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        expect(supplier.locator('#message')).to_have_value('')
        assert not supplier.evaluate("adapter('SEND',{jobId:'forbidden',expectedUrl:location.href,profile:{chat:'body',editor:'#message',send:'#send'}})")['clicked']
        passed('inquiry is neither filled nor sent, even with explicit selectors')
        supplier.evaluate("document.body.innerHTML='<a href=/sendInquiry id=contact>Send inquiry</a>'")
        assert not contact_call("adapter('CONTACT',{profile:{contact:'#contact'}})")['ok']
        passed('explicit contact selector cannot click Send inquiry')
        supplier.evaluate("document.body.innerHTML='<textarea id=message></textarea><button id=send>Send</button>'")
        assert supplier.evaluate("adapter('PROBE')")['kind']!='chat'
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        passed('generic textarea and Send without conversation context are not a chat')

        captures={
            'https://test.en.made-in-china.com/captured': 'captured-mic-store.html',
            'https://test.en.alibaba.com/captured': 'captured-alibaba.html',
            'https://test.1688.com/captured': 'captured-1688-store.html',
            'https://air.1688.com/captured-shell': 'captured-1688-shell.html',
            'https://example-shop.1688.com/': '1688-metadata.html',
            'https://air.1688.com/rich-composer': 'rich-composer-1688.html',
            'https://air.1688.com/delayed-composer': 'rich-composer-1688.html',
            'https://chat-test.1688.com/rich-composer': 'rich-composer-1688.html',
            'https://air.1688.com/composer-frame': '1688-composer-frame.html',
        }
        # Every iframe/resource request is intercepted locally. No vendor request
        # or challenge is fetched, even though the fixtures preserve frame URLs.
        delayed_resources=[]
        def capture_route(route):
            if route.request.url.endswith("/slow-ready-image"):
                delayed_resources.append(route)
                return
            filename=captures.get(route.request.url)
            if filename:
                body=(ROOT/'tests/fixtures'/filename).read_text()
                if route.request.url.endswith('/delayed-composer'):
                    body=body.replace('</body>', '<img hidden src="/slow-ready-image"></body>')
                route.fulfill(body=body,content_type='text/html')
            else:
                route.fulfill(body='<!doctype html><title>Empty frame fixture</title>',content_type='text/html')
        context.route('**/*',capture_route)
        supplier.goto('https://test.en.made-in-china.com/captured')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        result=supplier.evaluate("adapter('PROBE')")
        assert result['kind']=='contact' and result['diagnostics']['adapter']=='mic'
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('headerClicks')==0 and supplier.evaluate('otherClicks')==1
        expect(supplier.locator('#inquiryContent')).to_have_value('')
        assert supplier.evaluate('inquiryClicks')==0
        passed('MIC reference order: first visible J-open-tm is selected; inquiry remains untouched')
        supplier.evaluate("document.querySelector('#inqueryForm').removeAttribute('action')")
        assert not supplier.evaluate("adapter('PREPARE',{profile:{chat:'body',editor:'#inquiryContent',send:'#inquirySend'}})")['ok']
        expect(supplier.locator('#inquiryContent')).to_have_value('')
        passed('MIC inqueryForm and inquiry markers stay blocked even when action URL is absent')
        supplier.evaluate("document.querySelector('.shop-sign__btns-wrap').insertAdjacentHTML('beforeend','<a class=J-open-tm href=javascript:;>Chat Now</a>')")
        assert supplier.evaluate("adapter('PROBE')")['kind']=='contact'
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('otherClicks')==2 and supplier.evaluate('headerClicks')==0
        passed('repeated MIC chat controls follow the reference selector/DOM order without stopping on duplicate hooks')

        supplier.goto('https://test.en.alibaba.com/captured')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        assert supplier.evaluate("adapter('PROBE')")['kind']=='chat'
        assert supplier.evaluate("adapter('PREPARE')")['ok']
        expect(supplier.locator('textarea.send-textarea')).to_have_value('Hello, WeChat: demo')
        assert supplier.evaluate("adapter('SEND',{expectedUrl:location.href})")['clicked']
        assert supplier.evaluate('sendClicks')==1
        assert supplier.locator('.send-tool-button').evaluate("node=>node.classList.contains('gray')")
        passed('Alibaba gray button without disabled remains clickable')
        supplier.evaluate("document.querySelector('textarea.send-textarea').value='';document.querySelector('.send-tool-button').disabled=true")
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        supplier.evaluate("document.querySelector('.send-tool-button').disabled=false;document.querySelector('.send-tool-button').setAttribute('aria-disabled','true')")
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        passed('native disabled and aria-disabled still prevent Alibaba sending')
        supplier.evaluate("document.querySelector('.send-box-container').remove()")
        assert supplier.evaluate("adapter('PROBE')")['kind']=='contact'
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('headerClicks')==1 and supplier.evaluate('toolbarClicks')==0
        supplier.evaluate("document.querySelector('.shop-sign-container').remove()")
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('toolbarClicks')==1
        passed('captured Alibaba contacts: Russian storefront action first, Chat Now! toolbar fallback')

        supplier.goto('https://test.1688.com/captured')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        assert 'CAPTCHA' in supplier.evaluate("adapter('PROBE')")['obstacle']
        assert not contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('contactClicks')==0
        initial_overlay=supplier.evaluate("adapter('PROBE')")
        assert initial_overlay['canDismiss']
        supplier.evaluate("window.savedModal=document.querySelector('.J_MIDDLEWARE_FRAME_WIDGET')")
        assert supplier.evaluate("key=>adapter('DISMISS_CHALLENGE',{dismissKey:key})",initial_overlay['dismissKey'])['closed']
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('contactClicks')==1
        passed('1688 native close control dismisses the known overlay before customer-service click')
        for appearance in range(3):
            supplier.evaluate('document.body.append(savedModal)')
            repeated=supplier.evaluate("adapter('PROBE')")
            assert repeated['canDismiss'] and repeated['dismissKey']!=initial_overlay['dismissKey']
            assert not supplier.evaluate("key=>adapter('DISMISS_CHALLENGE',{dismissKey:key})",initial_overlay['dismissKey'])['clicked']
            assert supplier.evaluate("key=>adapter('DISMISS_CHALLENGE',{dismissKey:key})",repeated['dismissKey'])['closed']
            assert not supplier.evaluate("adapter('PROBE')").get('obstacle')
        passed('reused 1688 modal closes on every reappearance and rejects stale close commands')
        supplier.evaluate("document.body.innerHTML='<div id=application><div class=service-entry><span>客服</span></div></div>';window.serviceClicks=0;document.querySelector('.service-entry').addEventListener('click',()=>serviceClicks++)")
        assert contact_call("adapter('CONTACT')")['ok']
        assert supplier.evaluate('serviceClicks')==1
        passed('1688 customer-service text works without a specific id and preserves event bubbling')

        supplier.goto('https://air.1688.com/captured-shell')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        result=supplier.evaluate("adapter('PROBE')")
        assert result['kind']=='waiting' and result['diagnostics']['chatFrames']==1
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        passed('captured 1688 shell: waits for real iframe editor and ignores feedback textarea')

        supplier.goto('https://example-shop.1688.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        metadata=supplier.evaluate("adapter('PROBE')")
        assert metadata['kind']=='contact' and 'touid=cnalichn' in metadata['chatRoute']
        ready=supplier.evaluate("adapterMessages.find(message=>message.type==='PAGE_READY')")
        assert ready['chatRoute']==metadata['chatRoute']
        assert '示例供应商甲' in supplier.evaluate("decodeURIComponent(adapterMessages.find(message=>message.type==='PAGE_READY').chatRoute)")
        passed('1688 inline seller metadata is exposed to the worker before any contact-button click')
        lookup=supplier.evaluate('OpenLinks1688.lookup(location.href)')
        assert lookup['url']==metadata['chatRoute']
        assert not supplier.evaluate("adapterMessages.some(message=>message.type==='CONTACT')")
        passed('background-style HTML lookup resolves a previously unknown 1688 seller without opening its chat button')

        supplier.goto('https://test.en.alibaba.com/captured')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        supplier.evaluate("""document.body.innerHTML='<div class=dm-compose><textarea class=dm-draft></textarea><button class=dm-submit><span>Deliver</span></button></div><button class=vendor-contact>Talk to our team</button>';window.unknownClicks=0;window.earlyPointerDown=0;document.querySelector('.dm-submit').addEventListener('pointerdown',()=>earlyPointerDown++);document.querySelector('.dm-submit').addEventListener('click',()=>unknownClicks++);document.querySelector('.dm-draft').value='Private draft must not be recorded';""")
        assert supplier.evaluate("adapter('PICK_START',{field:'editor',token:'editor-lesson'})")['ok']
        supplier.locator('textarea.dm-draft').click()
        supplier.wait_for_function('lessonMessages.length===1')
        editor_pick=supplier.evaluate('lessonMessages[0]')
        assert editor_pick['field']=='editor' and 'Private draft' not in str(editor_pick)
        supplier.evaluate("window.lessonRules=[OpenLinksRules.makeRule(lessonMessages[0],location.href,'alibaba')]")
        assert supplier.evaluate("adapter('PICK_START',{field:'send',token:'send-lesson',profile:{rules:lessonRules}})")['ok']
        supplier.locator('button.dm-submit span').click()
        supplier.wait_for_function('lessonMessages.length===2')
        send_pick=supplier.evaluate('lessonMessages[1]')
        assert send_pick['field']=='send' and send_pick['chatSelector']
        assert supplier.evaluate('unknownClicks')==0
        assert supplier.evaluate('earlyPointerDown')==0
        passed('trusted teaching click records editor/send and inferred chat root before any native send')
        supplier.evaluate("""const p=lessonMessages[1];lessonRules.push(OpenLinksRules.makeRule(p,location.href,'alibaba'),OpenLinksRules.makeRule({field:'chat',selector:p.chatSelector},location.href,'alibaba'));document.querySelector('.dm-draft').value='';""")
        assert supplier.evaluate("adapter('PREPARE',{profile:{rules:lessonRules}})")['ok']
        assert supplier.evaluate("adapter('SEND',{jobId:'learned-send',expectedUrl:location.href,profile:{rules:lessonRules}})")['clicked']
        assert supplier.evaluate('unknownClicks')==1
        passed('learned nonstandard composer is filled and clicked exactly once on replay')
        assert supplier.evaluate("adapter('PICK_START',{field:'contact',token:'cancel-lesson'})")['ok']
        supplier.locator('[data-openlinks-picker] button').click()
        assert supplier.evaluate("document.querySelector('[data-openlinks-picker]')===null")
        assert supplier.evaluate('lessonMessages.length')==2
        passed('cancel teaching removes the overlay and records no unrelated click')
        persisted_rules=supplier.evaluate('JSON.stringify(lessonRules)')
        fresh=browser.new_context()
        fresh.on('page',lambda page:page.on('pageerror',lambda error:errors.append(str(error))))
        # A new browser context has no picker or page globals from the lesson.
        # Restore serialized rules to a different supplier and changed root layout.
        fresh.route('**/*',lambda route:route.fulfill(body='<div class=new-wrapper><section><textarea class=dm-draft></textarea></section><button class=dm-submit>Deliver</button></div><script>window.replayed=0;document.querySelector("button").onclick=()=>replayed++;</script>',content_type='text/html'))
        restored=fresh.new_page()
        restored.add_init_script(adapter_transport)
        restored.goto('https://second.en.alibaba.com/')
        restored.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        restored.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        restored.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        restored.add_script_tag(path=str(ROOT/'extension/content.js'))
        restored.evaluate("rules=>{window.restoredRules=JSON.parse(rules);restoredRules.push(OpenLinksRules.makeRule({field:'send',selector:'.missing-template'},location.href,'alibaba'));}",persisted_rules)
        replay=restored.evaluate("adapter('PROBE',{profile:{rules:restoredRules}})")
        assert replay['kind']=='chat' and replay['diagnostics']['learning']['derivedRoot']
        assert any('.missing-template' in reason for reason in replay['diagnostics']['learning']['rejected'])
        assert restored.evaluate("adapter('PREPARE',{profile:{rules:restoredRules}})")['ok']
        assert restored.evaluate("adapter('SEND',{expectedUrl:location.href,profile:{rules:restoredRules}})")['clicked']
        assert restored.evaluate('replayed')==1
        passed('serialized lessons work in a fresh context on another Alibaba shop with changed chat wrapper and alternative templates')
        fresh.close()
        supplier.goto('https://first-named.1688.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        supplier.evaluate("document.body.innerHTML='<button class=vendor-contact>Talk to our team</button>'")
        assert supplier.evaluate("adapter('PICK_START',{field:'contact',token:'named-shop'})")['ok']
        supplier.locator('.vendor-contact').click()
        supplier.wait_for_function('lessonMessages.length===1')
        named_rule=supplier.evaluate("JSON.stringify(OpenLinksRules.makeRule(lessonMessages[0],location.href,'1688'))")
        supplier.goto('https://second-named.1688.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/routes-1688-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        supplier.evaluate("rule=>{window.namedRule=JSON.parse(rule);document.body.innerHTML='<button class=vendor-contact>Talk to our team</button>';window.namedClicks=0;document.querySelector('button').onclick=()=>namedClicks++;}",named_rule)
        result=supplier.evaluate("adapter('PROBE',{profile:{rules:[namedRule]}})")
        assert result['kind']=='contact' and result['diagnostics']['learning']['applied']
        assert contact_call("adapter('CONTACT',{profile:{rules:[namedRule]}})")['ok']
        assert supplier.evaluate('namedClicks')==1
        passed('trusted 1688 contact lesson transfers between different named subdomains after page reload')
        # The user's draft is visibly present but Send remains inactive until
        # native input updates the editor's model. Include Unicode, markup and a newline.
        supplier.goto('https://air.1688.com/delayed-composer',wait_until='domcontentloaded')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        assert supplier.evaluate('document.readyState')!='complete'
        assert supplier.evaluate("adapter('PREPARE')")['waiting']
        assert supplier.locator('pre').inner_text()=='' and supplier.evaluate('mouseClicks')==0
        for resource in delayed_resources:
            resource.fulfill(status=204,body='')
        supplier.wait_for_load_state('load')
        assert supplier.evaluate("adapter('CHAT_READY')")['waiting']
        assert supplier.evaluate("adapter('PREPARE')")['waiting']
        supplier.wait_for_timeout(5100)
        assert supplier.evaluate("adapter('CHAT_READY')")['ok']
        passed('1688 input waits for a genuinely pending page load, then a stable editor for five seconds')
        rich_message='Hello <b>供应商</b>!\nWeChat: test & demo'
        # Run the production keyboard routine through a real Chromium CDP
        # session. Only chrome.debugger transport is replaced; input is trusted.
        cdp=context.new_cdp_session(supplier)
        native_source=(ROOT/'extension/lib/native-input.js').read_text().replace('export async function activate1688(', 'window.activate1688 = async function (')
        native_transport="""() => {
          window.attaches=0;window.detaches=0;window.nativeCommands=[];window.nativeKeys=[];window.nativeSpaceCodes=[];
          chrome.debugger={attach:async()=>{attaches++},detach:async()=>{detaches++},sendCommand:(target,method,params)=>new Promise(resolve=>{nativeKeys.push(params);nativeCommands.push({method,params,resolve})})};
          window.prepareNative=async(message='Hello, WeChat: demo')=>{
            const draft=await adapter('PREPARE',{message,reactivate:window.reactivate===true});window.reactivate=false;
            if(!draft.needsNativeInput)return draft;
            await activate1688({platform:'1688',tabId:1,frameId:0,inputUrl:draft.pageUrl},(job,frame,type,extra)=>{window.beforeNativeCheck?.(extra.stage);return adapter(type,{message,...extra})});
            return adapter('PREPARE',{message});
          };
        }"""
        supplier.add_script_tag(content=native_source)
        supplier.evaluate(native_transport)
        def prepare_native(message='Hello, WeChat: demo',target=supplier,expect_error=False):
            # Dispatch outside a pending Runtime.evaluate/exposed binding on
            # the input target, just as the extension worker does in Chrome.
            target.evaluate("message=>{window.nativeResult=null;prepareNative(message).then(value=>{nativeResult={value}},error=>{nativeResult={error:error.message,inputFailure:error.inputFailure,inputStage:error.inputStage}});}",message)
            deadline=time.monotonic()+10
            while time.monotonic()<deadline:
                command=target.evaluate('nativeCommands[0] && {method:nativeCommands[0].method,params:nativeCommands[0].params}')
                if command:
                    reply=cdp.send(command['method'],command['params'])
                    if command['method']=='Input.dispatchKeyEvent' and command['params']['type']=='char':
                        target.evaluate("()=>{nativeSpaceCodes.push(document.querySelector('pre').textContent.slice(-1).charCodeAt(0));}")
                    target.evaluate('reply=>{nativeCommands.shift().resolve(reply);}',reply)
                result=target.evaluate('nativeResult')
                if result:
                    if expect_error:
                        assert 'error' in result,result
                        return result
                    assert 'error' not in result,result
                    return result['value']
                supplier.wait_for_timeout(10)
            raise AssertionError('Chrome keyboard preparation timed out')
        supplier.evaluate("message=>{document.querySelector('pre').textContent=message;}",rich_message)
        assert supplier.evaluate('model')==''
        prepared=supplier.evaluate("message=>adapter('PREPARE',{message})",rich_message)
        assert not prepared['ok'] and prepared['needsNativeInput'], prepared
        assert supplier.locator('button').get_attribute('aria-disabled')=='true'
        prepared=prepare_native(rich_message)
        assert prepared['ok'] and prepared['inputMethod']=='chrome-mouse-keyboard', prepared
        assert supplier.evaluate('mouseClicks')==1
        assert supplier.locator('[data-openlinks-input]').count()==0
        assert supplier.evaluate('lastTrustedKey')=='Backspace'
        assert supplier.evaluate('nativeSpaceCodes')==[0xA0]
        assert supplier.locator('pre').inner_text()==rich_message
        assert supplier.evaluate('attaches')==supplier.evaluate('detaches')==1
        assert supplier.evaluate('model')==rich_message and supplier.evaluate('nativeEdits')>0
        assert supplier.locator('pre b').count()==0
        edits=supplier.evaluate('nativeEdits')
        assert supplier.evaluate("message=>adapter('PREPARE',{message})",rich_message)['ok']
        assert supplier.evaluate('nativeEdits')==edits
        sent=supplier.evaluate("message=>adapter('SEND',{message,expectedUrl:location.href})",rich_message)
        assert sent['clicked'] and sent['accepted'] and supplier.evaluate('sendClicks')==1
        expect(supplier.locator('#history .outgoing')).to_have_text(rich_message)
        assert not supplier.evaluate("message=>adapter('SEND',{message,expectedUrl:location.href})",rich_message)['clicked']
        passed('1688 requires a trusted mouse click then Space/Backspace: Chrome activates Send, preserves text, detaches and removes its marker')
        # A handler can be installed late on the same editor. A retry must
        # re-run trusted input even though the old preparation marker survived.
        supplier.evaluate("()=>{document.querySelector('pre').textContent='';window.acceptKeyboard=false;}")
        assert not prepare_native()['ok']
        clicks_before=supplier.evaluate('mouseClicks')
        supplier.evaluate('acceptKeyboard=true')
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        supplier.evaluate('window.reactivate=true')
        assert prepare_native()['ok']
        assert supplier.evaluate('mouseClicks')==clicks_before+1
        assert supplier.evaluate("adapter('SEND',{attemptId:'late-handler',expectedUrl:location.href})")['accepted']
        passed('retry reactivates the same 1688 editor after late keyboard handlers become ready')
        # The attached capture has a one-character draft followed by &nbsp;.
        # Test that shape and a composer that preserves ordinary spaces.
        for whitespace,message in [('pre-line','1'),('pre','2')]:
            supplier.evaluate("([ws,message])=>{let e=document.querySelector('pre');e.style.whiteSpace=ws;e.textContent=message+'\\u00a0';}",[whitespace,message])
            prepared=prepare_native(message)
            assert prepared['ok'],prepared
            assert supplier.locator('pre').inner_text()==message
            assert supplier.evaluate('model')==message
            assert supplier.evaluate('nativeSpaceCodes.at(-1)')==(0xA0 if whitespace=='pre-line' else 0x20)
            assert supplier.evaluate("message=>adapter('SEND',{message,attemptId:'space-'+getComputedStyle(document.querySelector('pre')).whiteSpace,expectedUrl:location.href})",message)['accepted']
        passed('captured 1688 NBSP draft recovers; both NBSP and ordinary native spaces complete without altering the message')

        # Validation runs after the click and both keys. A changed message still
        # cannot be sent, but temporary browser whitespace is not inspected.
        sends=supplier.evaluate('sendClicks')
        for replacement in ['12','2']:
            supplier.evaluate("text=>{let e=document.querySelector('pre');e.textContent='';window.nativeKeys=[];window.beforeNativeCheck=stage=>{if(stage==='done')e.textContent=text;};}",replacement)
            failed=prepare_native('1',expect_error=True)
            assert failed['inputFailure']=='editor' and failed['inputStage']=='done',failed
            assert supplier.evaluate("nativeKeys.some(k=>k.key==='Backspace')")
            assert supplier.evaluate('sendClicks')==sends
            assert supplier.evaluate('attaches')==supplier.evaluate('detaches')
        supplier.evaluate("()=>{window.beforeNativeCheck=null;document.querySelector('pre').textContent='';}")
        passed('final message validation stops a changed draft after the input sequence, without sending and with debugger released')
        supplier.evaluate('ignoreSend=true')
        assert prepare_native()['ok']
        ignored=supplier.evaluate("adapter('SEND',{attemptId:'ignored',expectedUrl:location.href})")
        assert ignored['clicked'] and not ignored['accepted']
        assert supplier.locator('pre').inner_text()=='Hello, WeChat: demo'
        assert supplier.evaluate('sendClicks')==sends+1
        assert not supplier.evaluate("adapter('SEND',{attemptId:'ignored',expectedUrl:location.href})")['clicked']
        passed('1688 ignored click is reported as unaccepted with draft intact and no automatic retry')
        supplier.evaluate("()=>{let e=document.querySelector('pre');e.innerHTML='<b>Old text</b><div>Second line</div>';window.replacementInputs=[];e.addEventListener('input',()=>replacementInputs.push(e.innerText));}")
        replaced=supplier.evaluate("adapter('PREPARE')")
        assert replaced['needsNativeInput'] and replaced['draftReplaced']
        assert supplier.locator('pre').inner_text()=='Hello, WeChat: demo'
        cleared=supplier.evaluate('replacementInputs')
        assert cleared[0].strip()=='' and '' in cleared,cleared
        assert prepare_native()['ok']
        assert supplier.evaluate('model')=='Hello, WeChat: demo'
        assert supplier.evaluate('sendClicks')==sends+1
        passed('1688 clears a multiline rich draft, activates the new message with native keys and does not send during preparation')

        # Closure marker belongs to this document only and is invalidated by
        # actual user interaction. A page cannot schedule another tab's close.
        assert supplier.evaluate("adapter('ARM_CLOSE',{token:'close-1',expectedUrl:location.href,at:Date.now()+10000})")['ok']
        assert supplier.evaluate("adapter('CAN_CLOSE',{token:'close-1'})")['ok']
        assert not supplier.evaluate("adapter('CAN_CLOSE',{token:'other'})")['ok']
        supplier.clock.install()
        supplier.clock.pause_at(int(time.time()*1000)+1000)
        supplier.evaluate("adapterMessages.length=0")
        assert supplier.evaluate("adapter('ARM_CLOSE',{token:'timed-close',expectedUrl:location.href,at:Date.now()+10000})")['ok']
        supplier.clock.run_for(9999)
        assert not supplier.evaluate("adapterMessages.some(m=>m.type==='CHAT_CLOSE_DUE')")
        supplier.clock.run_for(1)
        assert supplier.evaluate("adapterMessages.filter(m=>m.type==='CHAT_CLOSE_DUE').map(m=>m.token)")==['timed-close']
        passed('content timer wakes the worker after 10 seconds, never before the deadline')
        supplier.clock.resume()
        supplier.locator('pre').click()
        assert not supplier.evaluate("adapter('CAN_CLOSE',{token:'timed-close'})")['ok']
        passed('a trusted user interaction cancels the document-specific auto-close marker')

        supplier.goto('https://air.1688.com/composer-frame')
        inner=supplier.frame(url='https://chat-test.1688.com/rich-composer')
        assert inner
        inner.add_script_tag(path=str(ROOT/'extension/content.js'))
        inner.add_script_tag(content=native_source)
        inner.evaluate(native_transport)
        assert inner.evaluate("adapter('CHAT_READY')")['waiting']
        supplier.wait_for_timeout(5100)
        supplier.locator('#outside').fill('Keep this unrelated draft')
        prepared=prepare_native(rich_message,target=inner)
        assert prepared['ok'] and inner.evaluate('lastTrustedKey')=='Backspace'
        assert inner.evaluate('mouseClicks')==1
        assert inner.locator('pre').inner_text()==rich_message
        assert supplier.locator('#outside').input_value()=='Keep this unrelated draft'
        assert inner.evaluate('model')==rich_message
        passed('real keyboard input targets the 1688 iframe composer while preserving a focused field on its parent page')
        cdp.detach()

        # A MIC lesson describes the minimized chat control, not the initial
        # storefront action. It must transfer to another shop and survive reload.
        supplier.goto('https://first.en.made-in-china.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        supplier.evaluate("document.body.innerHTML='<div class=webtm-min-box aria-label=tm data-testid=min-first style=width:100px;height:40px>Supplier One</div>'")
        assert supplier.evaluate("adapter('PICK_START',{field:'contact',token:'mic-minimized'})")['ok']
        supplier.locator('[aria-label=tm]').click()
        supplier.wait_for_function('lessonMessages.length===1')
        mic_rule=supplier.evaluate("JSON.stringify(OpenLinksRules.makeRule(lessonMessages[0],location.href,'mic'))")
        assert 'Supplier One' not in mic_rule
        supplier.goto('https://second.en.made-in-china.com/')
        supplier.add_script_tag(path=str(ROOT/'extension/rules-shared.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        supplier.evaluate("rule=>{window.micRule=JSON.parse(rule);document.body.innerHTML='<div class=webtm-min-box aria-label=tm data-testid=min-second style=width:100px;height:40px>Supplier Two</div>';window.minimizedClicks=0;document.querySelector('[aria-label=tm]').onclick=()=>minimizedClicks++;}",mic_rule)
        result=supplier.evaluate("adapter('PROBE',{profile:{rules:[micRule]}})")
        assert result['kind']=='contact' and result['diagnostics']['learning']['applied']
        assert any('min-first' in reason for reason in result['diagnostics']['learning']['rejected'])
        assert contact_call("adapter('CONTACT',{profile:{rules:[micRule]}})")['ok']
        assert supplier.evaluate('minimizedClicks')==1
        # All selectors now match multiple controls, so the rule must not choose
        # an arbitrary supplier from a minimized-conversations list.
        supplier.evaluate("document.body.append(document.querySelector('[aria-label=tm]').cloneNode(true))")
        assert supplier.evaluate("adapter('PROBE',{profile:{rules:[micRule]}})")['kind']!='contact'
        passed('MIC minimized-chat lesson survives supplier/test-ID changes via saved alternatives; ambiguous recipients are rejected')
        supplier.goto('https://webim.made-in-china.com/super-tm/')
        supplier.evaluate("document.body.innerHTML='<section data-chat-root><div contenteditable=true class=webtm-input>Old MIC draft</div><button class=webtm-send-button>Send</button></section><textarea id=unrelated>Keep outside field</textarea>';window.clearStages=[];document.querySelector('.webtm-input').addEventListener('input',e=>clearStages.push(e.target.innerText));")
        supplier.add_script_tag(path=str(ROOT/'extension/supplier-chat.js'))
        supplier.add_script_tag(path=str(ROOT/'extension/content.js'))
        assert not supplier.evaluate("adapter('PREPARE')")['ok']
        assert supplier.locator('.webtm-input').inner_text()=='Old MIC draft'
        replaced=supplier.evaluate("adapter('PREPARE',{profile:{editor:'.webtm-input'}})")
        assert replaced['ok'] and replaced['draftReplaced']
        assert supplier.locator('.webtm-input').inner_text()=='Hello, WeChat: demo'
        assert supplier.locator('#unrelated').input_value()=='Keep outside field'
        cleared=supplier.evaluate('clearStages')
        assert cleared[0].strip()=='' and '' in cleared,cleared
        passed('Made-in-China clears and replaces the selected rich editor while leaving unrelated fields intact')


        ui.get_by_role('button',name='Очередь поставщиков').click()
        row=ui.locator('#queue-body tr').filter(has_text='jcnyjxsb.en.alibaba.com')
        row.get_by_role('button',name='Удалить',exact=True).click()
        expect(ui.locator('#queue-count')).to_have_text('11')
        expect(row).to_have_count(0)
        ui.locator('#supplier-input').fill('https://jcnyjxsb.en.alibaba.com/')
        ui.locator('#add').click()
        expect(ui.locator('#queue-count')).to_have_text('12')
        expect(ui.locator('#import-result')).to_contain_text('Добавлено: 1')
        expect(ui.locator('#queue-body tr').last).to_contain_text('jcnyjxsb.en.alibaba.com')
        passed('individual delete button removes one supplier and reimport places it at the end of the queue')
        assert not errors, errors
        passed('no browser JavaScript errors')
        browser.close()
finally:
    server.shutdown()
print(f'{checks} browser checks passed. All contacts used local fixtures; no seller was contacted.')
