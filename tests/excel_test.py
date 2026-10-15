"""Validate the browser's OOXML exporter with an independent spreadsheet reader."""
import io
import os
from pathlib import Path
import subprocess
from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent.parent
script=r'''
import {initialState} from './extension/lib/core.js';
import {reportExcel} from './extension/lib/excel.js';
const state=initialState(),at=Date.parse('2026-10-16T00:15:30.000Z');
state.events=[
 {at,jobId:'old',url:'https://shop.1688.com/?a=1&b=2',type:'attempted',detail:'Clicked once'},
 {at:at+1000,jobId:'old',url:'https://shop.1688.com/?a=1&b=2',type:'remove_job',detail:'Removed'},
 {at:at+2000,jobId:'new',url:'https://shop.1688.com/?a=1&b=2',type:'failed',errorCount:3,detail:'=HYPERLINK("https://example.test", "Formula must remain text")'},
 {at:at+3000,jobId:'unknown',url:'https://shop.en.alibaba.com/',type:'attempt_unknown',detail:'A & B <script>\nПример 中文 😀\u0001'.replaceAll('\\n','\n')},
];
if(process.argv[1]==='empty')state.events=[];
process.stdout.write(reportExcel(state));
'''
def workbook(mode='data'):
    result=subprocess.run(['node','--input-type=module','-e',script,mode],cwd=ROOT,env={**os.environ,'TZ':'Asia/Irkutsk'},check=True,capture_output=True)
    return result.stdout,load_workbook(io.BytesIO(result.stdout))

raw,book=workbook()
with ZipFile(io.BytesIO(raw)) as archive:
    assert archive.testzip() is None
    assert archive.comment==b'OpenLinks report.xlsx'
    for name in archive.namelist(): ElementTree.fromstring(archive.read(name))
print('PASS valid ZIP checksums, XML parts, relationships and independent XLSX loading')
assert book.sheetnames==['Результаты','События']
results,events=book.worksheets
assert results.freeze_panes=='A2' and events.freeze_panes=='A2'
assert results.auto_filter.ref=='A1:I4' and events.auto_filter.ref=='A1:G5'
print('PASS two worksheets with frozen headers and filters')
rows={row[8].value:row for row in list(results.rows)[1:]}
assert len(rows)==3 and rows['old'][3].value=='Кнопка отправки нажата'
assert rows['new'][4].value==3 and rows['unknown'][3].value=='Нужна проверка — без повтора'
print('PASS removed results, repeated URLs, error counts and unknown attempts retain distinct outcomes')
assert rows['new'][5].value==datetime(2026,10,16,8,15,32)
assert rows['new'][5].is_date and rows['new'][6].value=='Asia/Irkutsk (UTC+08:00)'
assert events['A2'].value==datetime(2026,10,16,8,15,30)
print('PASS real Excel datetime values match the computer wall clock and include the timezone')
assert rows['new'][7].value.startswith('=HYPERLINK') and rows['new'][7].data_type=='s'
assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
assert rows['new'][2].hyperlink.target=='https://shop.1688.com/?a=1&b=2'
print('PASS text cannot become an Excel formula; supplier links remain clickable')
assert 'Пример 中文 😀' in rows['unknown'][7].value and '<script>' in rows['unknown'][7].value and '\n' in rows['unknown'][7].value and '\x01' not in rows['unknown'][7].value
print('PASS Unicode, markup-like strings and XML escaping survive the spreadsheet round trip')
_,empty=workbook('empty')
assert all(sheet.max_row==1 for sheet in empty)
print('PASS an empty queue creates a valid workbook with headers')
(ROOT/'test-results').mkdir(exist_ok=True)
(ROOT/'test-results'/'report-example.xlsx').write_bytes(raw)
print('7 Excel checks passed.')
